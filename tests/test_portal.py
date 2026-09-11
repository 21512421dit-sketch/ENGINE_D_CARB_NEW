import io
import json
import os
import re
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from app import create_app, db
from app.models import Delivery, EngineCentre, EngineCentreContact, Lead, Submission, User

os.environ['ADMIN_EMAIL'] = 'admin@test.local'
os.environ['ADMIN_PASSWORD'] = 'TestPass123!'


def make_app(tmp_path):
    return create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': 'sqlite:///' + str(tmp_path / 'portal.db'), 'SECRET_KEY': 'test'})


def login(client, email, password):
    page = client.get('/admin/login')
    token = re.search(b'name="csrf" value="([^"]+)', page.data).group(1).decode()
    return client.post('/admin/login', data={'csrf': token, 'email': email, 'password': password}, follow_redirects=False)


def engine_service(**changes):
    data = {'enquiryType': 'service', 'customerName': 'Engine Customer', 'servicePhone': '9876543210',
            'vehicleType': 'Car', 'vehicleBrand': 'Tata', 'vehicleModel': 'Nexon', 'passingYear': '2022',
            'fuelType': 'Diesel', 'engineCc': '1497', 'kilometres': '45000', 'area': 'Baner',
            'serviceCity': 'Pune', 'servicePin': '411045', 'selectedCentre': 'chikalthana-midc',
            'consent': True}
    return data | changes


def test_consent_engine_pricing_and_retention(tmp_path):
    app = make_app(tmp_path)
    client = app.test_client()
    public_page = client.get('/engine-d-carb')
    assert b'Nothing is uploaded automatically' not in public_page.data
    assert b'engine-integration.js' in public_page.data
    assert public_page.data.count(b'<details><summary>') == 7
    assert b'"@type": "FAQPage"' in public_page.data
    assert client.post('/api/engine-d-carb/quotations', json=engine_service(consent=False)).status_code == 400
    response = client.post('/api/engine-d-carb/quotations', json=engine_service())
    assert response.status_code == 200
    assert response.json['indicative_cost'] == 2058
    assert response.json['centre']['name'] == 'Chikalthana MIDC'
    assert response.json['messages']['centre_numbers'] == ['7727005151']
    assert 'Thank you Engine Customer' in response.json['messages']['customer']
    assert 'An estimated cost of ₹2,058' in response.json['messages']['centre']
    with app.app_context():
        item = Submission.query.one()
        assert item.site == 'engine_dcarb' and item.submitted_by_id is None
        assert 729 <= (item.expires_at.replace(tzinfo=timezone.utc) - item.created_at.replace(tzinfo=timezone.utc)).days <= 731


def test_engine_centres_are_public_and_admin_manages_multiple_numbers(tmp_path):
    app = make_app(tmp_path)
    client = app.test_client()
    centres = client.get('/api/engine-d-carb/centres')
    assert centres.status_code == 200
    assert [item['name'] for item in centres.json['centres']] == [
        'Chikalthana MIDC', 'Waluj MIDC', 'Beed By Pass', 'Garkheda'
    ]
    assert 'contacts' not in centres.json['centres'][0]
    assert client.post('/api/engine-d-carb/quotations', json=engine_service(selectedCentre='unknown')).status_code == 400

    login(client, 'admin@test.local', 'TestPass123!')
    page = client.get('/admin?site=engine_dcarb')
    csrf = re.search(b'name="csrf" value="([^"]+)', page.data).group(1).decode()
    with app.app_context():
        centre = EngineCentre.query.filter_by(key='waluj-midc').one()
        centre_id = centre.id
    added = client.post(f'/admin/engine-centres/{centre_id}/contacts', data={
        'csrf': csrf, 'contact_name': 'Workshop manager', 'phone': '9000012345'
    }, follow_redirects=False)
    assert added.status_code == 302
    with app.app_context():
        assert EngineCentreContact.query.filter_by(centre_id=centre_id).count() == 2
    quote = client.post('/api/engine-d-carb/quotations', json=engine_service(selectedCentre='waluj-midc'))
    assert quote.json['messages']['centre_numbers'] == ['8796344838', '9000012345']
    assert 'Siddhivinayak Wheel Alignment Center' in quote.json['messages']['customer']


def test_employee_empty_dropdowns_allow_manual_entry(tmp_path):
    app = make_app(tmp_path)
    client = app.test_client()
    script = client.get('/static/employee-portal.js')
    assert script.status_code == 200
    assert b"field.type === 'select' && selectOptions.length" in script.data
    assert b"Enter ${field.label.toLowerCase()} manually" in script.data


def test_admin_adds_employee_employee_machine_and_filtered_excel(tmp_path):
    app = make_app(tmp_path)
    client = app.test_client()
    assert login(client, 'admin@test.local', 'TestPass123!').headers['Location'].endswith('/admin')
    admin_page = client.get('/admin')
    csrf = re.search(b'name="csrf" value="([^"]+)', admin_page.data).group(1).decode()
    added = client.post('/admin/employees', data={'csrf': csrf, 'email': 'employee@example.com', 'password': 'Employee123!'}, follow_redirects=False)
    assert added.status_code == 302
    client.post('/admin/logout', data={'csrf': csrf})
    assert login(client, 'employee@example.com', 'Employee123!').headers['Location'].endswith('/portal')
    response = client.post('/api/engine-d-carb/quotations', json=engine_service(expertMachine='DCC PD10K'))
    assert response.status_code == 200 and response.json['machine'] == 'DCC PD10K'
    with app.app_context():
        item = Submission.query.one()
        assert item.submitted_by.email == 'employee@example.com'
    client.get('/admin/logout')
    client = app.test_client()
    login(client, 'admin@test.local', 'TestPass123!')
    admin_page = client.get('/admin?site=engine_dcarb')
    assert b'class="details-button"' in admin_page.data
    assert b'<dialog class="details-dialog"' in admin_page.data
    assert b'Engine capacity (CC)' in admin_page.data
    assert b'engineCc' not in admin_page.data
    exported = client.get('/admin/export.xlsx?site=engine_dcarb&columns=name&columns=employee&columns=result')
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.data))
    sheet = book.active
    assert sheet.title == 'Engine D-Carb'
    assert [cell.value for cell in sheet[1]] == ['Customer name', 'Submitted by', 'Quotation result']
    assert sheet.cell(2, 1).value == 'Engine Customer' and sheet.cell(2, 2).value == 'employee@example.com'
    result_text = sheet.cell(2, 3).value
    assert 'Status: Ready\nIndicative cost: ₹2,058\nRecommended machine: DCC PD10K' in result_text
    assert 'Centre — Name: Chikalthana MIDC' in result_text

    detailed = client.get('/admin/export.xlsx?site=engine_dcarb&columns=details')
    detail_sheet = load_workbook(io.BytesIO(detailed.data)).active
    detail_text = detail_sheet.cell(2, 1).value
    assert 'Engine capacity (CC): 1497' in detail_text
    assert 'consent' not in detail_text.lower()
    assert not detail_text.lstrip().startswith('{')


def test_expiry_purges_linked_customer_and_delivery_data(tmp_path):
    app = make_app(tmp_path)
    with app.app_context():
        lead = Lead(name='Expired Customer', phone='9000000000', form_json='{}', result_json='{}')
        db.session.add(lead); db.session.flush()
        db.session.add(Delivery(lead_id=lead.id, channel='sms', target='9000000000', status='sent'))
        db.session.add(Submission(site='batterywala', form_kind='quotation', name=lead.name, phone=lead.phone,
                                  form_json='{}', result_json='{}', consented=True, consent_version='test',
                                  lead_id=lead.id, expires_at=datetime.now(timezone.utc) - timedelta(days=1)))
        db.session.commit()
        from app.portal import purge_expired_submissions
        assert purge_expired_submissions() == 1
        assert Lead.query.count() == Delivery.query.count() == Submission.query.count() == 0


def test_admin_can_view_and_download_saved_quotation_pdf(tmp_path):
    from app import quotations
    app = make_app(tmp_path)
    client = app.test_client()
    lead = None
    with app.app_context():
        form = {'name': 'PDF Customer', 'phone': '9876543210', 'application': 'Passenger vehicle',
                'exchange_old_battery': 'no'}
        quote = {'number': 'BW-000001', 'date': '08/09/2026', 'options': [], 'status': 'pending_review',
                 'message': 'Price confirmation required.', 'note': quotations.NOTES['no'], 'token': 'secret'}
        lead = Lead(name=form['name'], phone=form['phone'], form_json=json.dumps(form),
                    result_json=json.dumps({'quotation': quote}))
        db.session.add(lead); db.session.flush()
        db.session.add(Submission(site='batterywala', form_kind='quotation', name=lead.name, phone=lead.phone,
                                  form_json=json.dumps(form), result_json=json.dumps({'quotation': quote}),
                                  consented=True, lead_id=lead.id, expires_at=datetime.now(timezone.utc) + timedelta(days=730)))
        db.session.commit(); lead_id = lead.id
    assert client.get(f'/admin/quotations/{lead_id}/pdf').status_code == 302
    login(client, 'admin@test.local', 'TestPass123!')
    page = client.get('/admin')
    assert f'/admin/quotations/{lead_id}/pdf'.encode() in page.data
    viewed = client.get(f'/admin/quotations/{lead_id}/pdf')
    assert viewed.status_code == 200 and viewed.mimetype == 'application/pdf'
    assert viewed.headers['Content-Disposition'].startswith('inline;')
    downloaded = client.get(f'/admin/quotations/{lead_id}/pdf?download=1')
    assert downloaded.headers['Content-Disposition'].startswith('attachment;')
