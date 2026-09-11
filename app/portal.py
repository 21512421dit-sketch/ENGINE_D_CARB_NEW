import calendar
import io
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path

from flask import Blueprint, Response, current_app, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from werkzeug.security import generate_password_hash

from . import db
from .models import Delivery, EngineCentre, EngineCentreContact, Lead, Submission, User


bp = Blueprint('portal', __name__)
CONSENT_VERSION = '2026-09'
SITES = {'batterywala': 'BatteryWala', 'engine_dcarb': 'Engine D-Carb'}
MACHINES = ('DCC AD6000', 'DCC PD6000', 'DCC PD10K', 'DCC PD15K')
ENGINE_CENTRES = (
    ('chikalthana-midc', 'Chikalthana MIDC', 'Care4Earth Enterprises, F-7/3, Chikalthana MIDC, Opp Flemingo Housing Society, Chhatrapati Sambhajinagar, Maharashtra.', '7727005151'),
    ('waluj-midc', 'Waluj MIDC', 'Siddhivinayak Wheel Alignment Center, Opp Tirupati Hospital, Mahaveer Chowk, Waluj MIDC, Chhatrapati Sambhajinagar, Maharashtra.', '8796344838'),
    ('beed-bypass', 'Beed By Pass', 'Master Garage, Opp Kamalnayan Bajaj Hospital, Beside LPG Fuel Pump, Beed By Pass Road, Chhatrapati Sambhajinagar, Maharashtra.', '9923778777'),
    ('garkheda', 'Garkheda', 'Fox 3D Wheel Alignment, Opp Sports Olympia Complex, Sut Girni Chowk, Garkheda, Chhatrapati Sambhajinagar, Maharashtra.', '9518537075'),
)
ENGINE_FAQS = (
    ('What is engine decarbonization?', 'It is a workshop service intended to address carbon deposits that accumulate during engine operation. Engine D-Carb uses a guided HHO cycle while the engine remains assembled.'),
    ('Is engine decarbonization good and safe for an engine?', 'Suitability depends on the vehicle and its condition. Engine D-Carb is designed as a non-invasive, chemical-free process and should be performed by a trained technician after basic vehicle checks.'),
    ('How much does it cost to decarbonize an engine?', 'Cost varies with vehicle class, fuel type, engine capacity and service location. Submit your vehicle details for a tentative quotation; the service team confirms the final price.'),
    ('Does decarbonizing increase mileage?', 'It may help recover efficiency lost to carbon build-up, but results vary with engine condition, maintenance, fuel quality and driving style. No fixed mileage improvement is guaranteed.'),
    ('How long does engine decarbonization take?', 'The Engine D-Carb service typically takes 45 to 90 minutes, depending on the vehicle, inspection and selected cycle.'),
    ('What signs may indicate carbon build-up?', 'Possible signs include rough idling, sluggish response, visible exhaust smoke or a change in fuel efficiency. These symptoms can have other causes, so a vehicle check remains important.'),
    ('How often should an engine be decarbonized?', 'There is no universal interval for every vehicle. Driving pattern, symptoms, mileage, maintenance history and the manufacturer’s guidance should inform the decision.'),
)
EXPORT_COLUMNS = {
    'created_at': 'Submitted at', 'form_kind': 'Form type', 'name': 'Customer name',
    'phone': 'Phone', 'email': 'Email', 'employee': 'Submitted by',
    'details': 'Submitted details', 'result': 'Quotation result',
    'consent': 'Consent', 'expires_at': 'Retention expiry',
}
FIELD_LABELS = {
    'area': 'Area', 'businessDetails': 'Business details', 'companyAddress': 'Company address',
    'customerName': 'Customer name', 'engineCc': 'Engine capacity (CC)', 'enquiryType': 'Enquiry type',
    'expertMachine': 'Recommended machine', 'fuelType': 'Fuel type', 'kilometres': 'Kilometres driven',
    'machineCity': 'City', 'machineEmail': 'Email address', 'machinePhone': 'Phone number',
    'machinePin': 'PIN code', 'passingYear': 'Registration year', 'representativeName': 'Representative name',
    'serviceCity': 'City', 'serviceDetails': 'Service details', 'servicePhone': 'Phone number',
    'servicePin': 'PIN code', 'vehicleBrand': 'Vehicle brand', 'vehicleModel': 'Vehicle model',
    'vehicleType': 'Vehicle type', 'indicative_cost': 'Indicative cost', 'machine': 'Recommended machine',
    'selectedCentre': 'Preferred service centre', 'message': 'Message', 'note': 'Note',
    'number': 'Quotation number', 'status': 'Status',
}
HIDDEN_DISPLAY_FIELDS = {'consent', 'token'}


def field_label(key):
    if key in FIELD_LABELS:
        return FIELD_LABELS[key]
    words = re.sub(r'([a-z0-9])([A-Z])', r'\1 \2', key).replace('_', ' ').strip()
    return words[:1].upper() + words[1:]


def display_value(value, key=''):
    if value is True:
        return 'Yes'
    if value is False:
        return 'No'
    if value in (None, ''):
        return 'Not provided'
    if key == 'indicative_cost':
        return f"₹{int(value):,}"
    if key == 'status':
        return str(value).replace('_', ' ').title()
    if isinstance(value, (list, tuple)):
        return ', '.join(str(item) for item in value) or 'Not provided'
    return str(value)


def display_fields(data, prefix=''):
    fields = []
    for key, value in data.items():
        if key in HIDDEN_DISPLAY_FIELDS:
            continue
        aliases = {'name': ('customerName', 'representativeName'), 'phone': ('servicePhone', 'machinePhone'),
                   'email': ('machineEmail',)}
        if key in aliases and any(alias in data for alias in aliases[key]):
            continue
        label = f'{prefix} — {field_label(key)}' if prefix else field_label(key)
        if isinstance(value, dict):
            fields.extend(display_fields(value, label))
        else:
            fields.append({'label': label, 'value': display_value(value, key)})
    return fields


def fields_as_text(fields):
    return '\n'.join(f"{field['label']}: {field['value']}" for field in fields)


def utcnow():
    return datetime.now(timezone.utc)


def retention_expiry(value=None):
    value = value or utcnow()
    year, month = value.year + 2, value.month
    day = min(value.day, calendar.monthrange(year, month)[1])
    return value.replace(year=year, month=month, day=day)


def ensure_engine_centres():
    for index, (key, name, address, default_phone) in enumerate(ENGINE_CENTRES, start=1):
        centre = EngineCentre.query.filter_by(key=key).first()
        if centre:
            centre.name, centre.address, centre.sort_order = name, address, index
            continue
        centre = EngineCentre(key=key, name=name, address=address, sort_order=index)
        centre.contacts.append(EngineCentreContact(contact_name='Primary centre contact', phone=default_phone))
        db.session.add(centre)
    db.session.commit()


def centre_dict(centre, include_contacts=False):
    data = {'key': centre.key, 'name': centre.name, 'address': centre.address}
    if include_contacts:
        data['contacts'] = [{'id': item.id, 'name': item.contact_name, 'phone': item.phone}
                            for item in centre.contacts]
    return data


def service_messages(form, result, centre):
    vehicle = f"{form['vehicleType']} — {form['vehicleBrand']} {form['vehicleModel']}, {form['passingYear']}, {form['fuelType']}, {form['engineCc']} CC, {int(form['kilometres']):,} km"
    numbers = [item.phone for item in centre.contacts]
    centre_details = f"{centre.name}\n{centre.address}"
    if numbers:
        centre_details += f"\nContact: {', '.join(numbers)}"
    customer = (
        f"Thank you {form['customerName']} for the enquiry of Engine D-Carb Service.\n\n"
        f"Vehicle details:\n{vehicle}\n\n"
        f"Your selected nearest centre details are as follows:\n{centre_details}"
    )
    owner = (
        f"{form['customerName']} is interested to avail the Engine D-Carb Service. "
        f"The contact number is {form['servicePhone']}.\n\n"
        f"Vehicle details:\n{vehicle}\n\n"
        f"An estimated cost of ₹{result['indicative_cost']:,} is given to the customer for availing Engine D-Carb service.\n\n"
        "Kindly contact the customer and book the appointment."
    )
    return {'customer': customer, 'centre': owner, 'centre_numbers': numbers, 'sender_number': '7727005151'}


def purge_expired_submissions():
    now = utcnow()
    expired = Submission.query.filter(Submission.expires_at <= now).all()
    if not expired:
        return 0
    lead_ids = {item.lead_id for item in expired if item.lead_id}
    if lead_ids:
        Delivery.query.filter(Delivery.lead_id.in_(lead_ids)).delete(synchronize_session=False)
        Lead.query.filter(Lead.id.in_(lead_ids)).delete(synchronize_session=False)
    for item in expired:
        db.session.delete(item)
    db.session.commit()
    return len(expired)


def record_submission(site, form_kind, form, result=None, lead_id=None, consented=True):
    if not consented:
        return None
    actor_id = current_user.id if current_user.is_authenticated and not current_user.is_admin else None
    submission = Submission(
        site=site, form_kind=form_kind, name=form.get('name') or form.get('customerName') or form.get('representativeName'),
        email=form.get('email') or form.get('machineEmail'), phone=form.get('phone') or form.get('servicePhone') or form.get('machinePhone'),
        form_json=json.dumps(form, ensure_ascii=False), result_json=json.dumps(result or {}, ensure_ascii=False),
        consented=True, consent_version=CONSENT_VERSION, submitted_by_id=actor_id,
        lead_id=lead_id, expires_at=retention_expiry(),
    )
    db.session.add(submission)
    return submission


def admin_required(fn):
    from functools import wraps

    @wraps(fn)
    @login_required
    def inner(*args, **kwargs):
        if not current_user.is_admin:
            return ('Forbidden', 403)
        return fn(*args, **kwargs)

    return inner


def engine_site_response():
    source = Path(current_app.root_path).parent / 'docs' / 'engine_dcarb_latest.html'
    html = source.read_text(encoding='utf-8')
    html = html.replace('Your information stays in your browser until you choose to send it.',
                        'Your information is stored for 24 months only after you accept the consent notice and submit.')
    html = html.replace('Submitting prepares the summary on this device. Nothing is uploaded automatically.',
                        'Submitting stores this request only after consent, then prepares the summary on this device.')
    html = html.replace('Complete the form, then continue on WhatsApp for tentative servicing cost and the nearest service-centre contact details.',
                        'Complete the form, choose your nearest centre and receive an estimated servicing cost.')
    html = html.replace('Continue on WhatsApp to receive tentative cost and nearby service-centre details.',
                        'WhatsApp messages are prepared for automatic delivery once the connection is configured.')
    faq_markup = ''.join(f'<details><summary>{question}</summary><p>{answer}</p></details>'
                         for question, answer in ENGINE_FAQS)
    html = re.sub(r'(<div class="faq-list reveal">).*?(</div></div></section>`;)',
                  rf'\1{faq_markup}\2', html, count=1, flags=re.DOTALL)
    faq_schema = {'@context': 'https://schema.org', '@type': 'FAQPage', 'mainEntity': [
        {'@type': 'Question', 'name': question,
         'acceptedAnswer': {'@type': 'Answer', 'text': answer}}
        for question, answer in ENGINE_FAQS]}
    integration = (
        '<link rel="icon" href="/static/images/batterywala-logo-original.png">'
        '<link rel="stylesheet" href="/static/engine-integration.css?v=20260909-1">'
        '<meta name="application-name" content="Engine D-Carb">'
        f'<script type="application/ld+json">{json.dumps(faq_schema, ensure_ascii=False)}</script></head>'
    )
    scripts = '<script src="/static/engine-integration.js?v=20260909-1"></script></body>'
    return Response(html.replace('</head>', integration).replace('</body>', scripts), mimetype='text/html')


@bp.get('/engine-d-carb')
def engine_site():
    return engine_site_response()


@bp.get('/api/engine-d-carb/centres')
def engine_centres():
    centres = EngineCentre.query.order_by(EngineCentre.sort_order, EngineCentre.id).all()
    return jsonify(centres=[centre_dict(item) for item in centres])


@bp.get('/portal')
@login_required
def workspace():
    if current_user.is_admin:
        return redirect(url_for('main.admin'))
    return render_template('employee.html', machines=MACHINES, consent_version=CONSENT_VERSION)


def valid_engine_payload(payload):
    if not isinstance(payload, dict) or len(payload) > 35:
        return None, 'Submit valid form data.'
    if any(not isinstance(value, (str, bool)) for value in payload.values()):
        return None, 'Submit valid form fields.'
    clean = {key: value.strip() if isinstance(value, str) else value for key, value in payload.items()}
    if any(len(value) > (2000 if key in {'serviceDetails', 'businessDetails', 'machineDetails', 'companyAddress'} else 255)
           for key, value in clean.items() if isinstance(value, str)):
        return None, 'One or more fields are too long.'
    return clean, None


@bp.post('/api/engine-d-carb/quotations')
def engine_quotation():
    form, error = valid_engine_payload(request.get_json(silent=True))
    if error:
        return jsonify(error=error), 400
    consented = form.get('consent') is True or form.get('consent') == 'true'
    if not consented:
        return jsonify(error='Consent is required before we can store your details and prepare the quotation.'), 400
    kind = form.get('enquiryType')
    if kind not in {'service', 'machine'}:
        return jsonify(error='Choose vehicle service or new machine enquiry.'), 400
    employee = current_user.is_authenticated and not current_user.is_admin
    selected_machine = form.get('expertMachine', '') if employee else ''
    if selected_machine and selected_machine not in MACHINES:
        return jsonify(error='Choose a valid Engine D-Carb machine.'), 400
    if kind == 'service':
        required = ('customerName', 'servicePhone', 'vehicleType', 'vehicleBrand', 'vehicleModel', 'passingYear',
                    'fuelType', 'engineCc', 'kilometres', 'area', 'serviceCity', 'servicePin', 'selectedCentre')
        if any(not form.get(key) for key in required):
            return jsonify(error='Complete all required vehicle service fields.'), 400
        if not re.fullmatch(r'[0-9]{10}', form['servicePhone']) or not re.fullmatch(r'[0-9]{6}', form['servicePin']):
            return jsonify(error='Enter a valid 10-digit phone number and 6-digit PIN code.'), 400
        try:
            cc, year, kilometres = int(form['engineCc']), int(form['passingYear']), int(form['kilometres'])
        except ValueError:
            return jsonify(error='Enter valid numbers for year, engine CC and kilometres.'), 400
        if not 999 <= cc <= 15000 or not 2010 <= year <= 2027 or kilometres < 0:
            return jsonify(error='Use 999–15,000 CC, a registration year from 2010–2027 and valid kilometres.'), 400
        factor = 1.1 if form['vehicleType'] == 'Car' else 1.5
        if form['vehicleType'] == 'Car' and form['fuelType'] == 'Diesel':
            factor = 1.25
        centre = EngineCentre.query.filter_by(key=form['selectedCentre']).first()
        if not centre:
            return jsonify(error='Choose a valid Engine D-Carb service centre.'), 400
        result = {'status': 'ready', 'indicative_cost': round(cc * factor * 1.1),
                  'machine': selected_machine or None, 'centre': centre_dict(centre)}
        result['messages'] = service_messages(form, result, centre)
        name, phone = form['customerName'], form['servicePhone']
    else:
        required = ('companyAddress', 'machineEmail', 'representativeName', 'machinePhone', 'businessDetails', 'machineCity', 'machinePin')
        if any(not form.get(key) for key in required):
            return jsonify(error='Complete all required machine enquiry fields.'), 400
        if not re.fullmatch(r'[0-9]{10}', form['machinePhone']) or not re.fullmatch(r'[0-9]{6}', form['machinePin']):
            return jsonify(error='Enter a valid 10-digit phone number and 6-digit PIN code.'), 400
        if not re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+', form['machineEmail']):
            return jsonify(error='Enter a valid email address.'), 400
        result = {'status': 'received', 'machine': selected_machine or None}
        name, phone = form['representativeName'], form['machinePhone']
    normalized = dict(form, name=name, phone=phone, email=form.get('machineEmail', ''))
    record_submission('engine_dcarb', kind, normalized, result, consented=True)
    db.session.commit()
    return jsonify(result)


@bp.post('/admin/engine-centres/<int:centre_id>/contacts')
@admin_required
def add_engine_centre_contact(centre_id):
    from .routes import valid_csrf

    if not valid_csrf():
        return ('Invalid CSRF', 400)
    centre = db.get_or_404(EngineCentre, centre_id)
    contact_name = request.form.get('contact_name', '').strip() or 'Centre contact'
    phone = re.sub(r'\D', '', request.form.get('phone', ''))
    if len(contact_name) > 120 or not re.fullmatch(r'[0-9]{10}', phone):
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='Enter a contact name and valid 10-digit number.'))
    if EngineCentreContact.query.filter_by(centre_id=centre.id, phone=phone).first():
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='That number is already assigned to this centre.'))
    db.session.add(EngineCentreContact(centre=centre, contact_name=contact_name, phone=phone))
    db.session.commit()
    return redirect(url_for('main.admin', site='engine_dcarb', centre_added='1'))


@bp.post('/admin/engine-centre-contacts/<int:contact_id>/delete')
@admin_required
def delete_engine_centre_contact(contact_id):
    from .routes import valid_csrf

    if not valid_csrf():
        return ('Invalid CSRF', 400)
    contact = db.session.get(EngineCentreContact, contact_id)
    if contact:
        db.session.delete(contact)
        db.session.commit()
    return redirect(url_for('main.admin', site='engine_dcarb', centre_deleted='1'))


@bp.post('/admin/employees')
@admin_required
def add_employee():
    from .routes import valid_csrf

    if not valid_csrf():
        return ('Invalid CSRF', 400)
    email = request.form.get('email', '').strip().lower()
    password = request.form.get('password', '')
    if not re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+', email) or len(password) < 8:
        return redirect(url_for('main.admin', employee_error='Use a valid email and a password of at least 8 characters.'))
    if User.query.filter_by(email=email).first():
        return redirect(url_for('main.admin', employee_error='That email already has an account.'))
    db.session.add(User(email=email, password_hash=generate_password_hash(password), is_admin=False))
    db.session.commit()
    return redirect(url_for('main.admin', employee_added='1'))


def filtered_submissions(args):
    site = args.get('site', 'batterywala')
    if site not in SITES:
        site = 'batterywala'
    query = Submission.query.filter_by(site=site, consented=True)
    kind = args.get('kind', '').strip()
    if kind:
        query = query.filter_by(form_kind=kind)
    employee_id = args.get('employee_id', '').strip()
    if employee_id.isdigit():
        query = query.filter_by(submitted_by_id=int(employee_id))
    for key, operator in (('from', 'from'), ('to', 'to')):
        raw = args.get(key, '')
        if not raw:
            continue
        try:
            value = datetime.fromisoformat(raw).replace(tzinfo=timezone.utc)
            if operator == 'to':
                value = value.replace(hour=23, minute=59, second=59)
            query = query.filter(Submission.created_at >= value) if operator == 'from' else query.filter(Submission.created_at <= value)
        except ValueError:
            pass
    return site, query.order_by(Submission.created_at.desc())


def submission_dict(item):
    form = json.loads(item.form_json or '{}')
    result = json.loads(item.result_json or '{}')
    return {
        'id': item.id, 'site': item.site, 'form_kind': item.form_kind, 'name': item.name or '',
        'phone': item.phone or '', 'email': item.email or '',
        'employee': item.submitted_by.email if item.submitted_by else 'Public website',
        'details': form, 'details_display': display_fields(form),
        'result': result, 'result_display': display_fields(result), 'consent': 'Granted',
        'created_at': item.created_at.isoformat(), 'expires_at': item.expires_at.isoformat(),
        'quotation_pdf_url': (url_for('quotations.admin_pdf', lead_id=item.lead_id)
                              if item.site == 'batterywala' and item.form_kind == 'quotation' and item.lead_id else None),
    }


@bp.get('/admin/submissions.json')
@admin_required
def submissions_json():
    purge_expired_submissions()
    site, query = filtered_submissions(request.args)
    rows = query.limit(200).all()
    return jsonify(site=site, total=query.count(), rows=[submission_dict(item) for item in rows])


@bp.get('/admin/export.xlsx')
@admin_required
def export_submissions():
    purge_expired_submissions()
    site, query = filtered_submissions(request.args)
    selected = [name for name in request.args.getlist('columns') if name in EXPORT_COLUMNS]
    if not selected:
        selected = list(EXPORT_COLUMNS)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SITES[site][:31]
    sheet.append([EXPORT_COLUMNS[name] for name in selected])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='102E3C')
        cell.alignment = Alignment(vertical='center')
    for item in query.all():
        data = submission_dict(item)
        values = []
        for name in selected:
            value = data[name]
            if name == 'details':
                value = fields_as_text(data['details_display'])
            elif name == 'result':
                value = fields_as_text(data['result_display'])
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(48, max(12, max(len(str(cell.value or '')) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
        for cell in column:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    stamp = utcnow().strftime('%Y-%m-%d')
    return send_file(output, as_attachment=True, download_name=f'{site}-submissions-{stamp}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
